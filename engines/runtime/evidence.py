"""
evidence.py — Evidence Collection Engine.

    collection = collect_evidence(execution_run, action_plan, test_cases, artifact_dir)

Turns what the Playwright Executor recorded into one evidence bundle per executed plan, linked to the BRD
acceptance criterion the test verifies. Facts only: no verdicts (Output Validation and Pass/Fail decide).

Per bundle:
  * inputs     — what was sent: each field value / message / HTTP body, flagged when it was filler test data;
  * observed   — what came back, normalised by output kind:
                 ui_output     text that appeared on the page after the actions (final page text minus the page
                               as first opened), plus validation / error messages shown;
                 reply         the chat reply;
                 documents     every downloaded file opened and read: text, headings, tables, pages, words
                               (DOCX, PDF, XLSX, PPTX, text formats; an unreadable file is recorded as such);
                 http_response status, timing, body excerpt;
  * artifacts  — every file with type, size and sha256 (screenshots, page text, downloads, extracted text);
  * observations — API calls (count, failures, slowest), console and page errors;
  * completeness — whether the output this kind of test needs was captured, and what is missing.
A manifest (evidence_manifest.json) with every bundle and file hash is written to the run's artifact folder,
so evidence can be re-checked later against the files on disk.
"""

import hashlib
import json
import re
import time
import zipfile
from pathlib import Path
from difflib import SequenceMatcher
from typing import Any, Dict, List, Optional, Sequence, Tuple

TEXT_LIMIT = 200_000
EXCERPT = 2_000
_MESSAGE = re.compile(r"\b(required|invalid|must|error|failed|failure|please (enter|fill|provide|select)|not allowed|"
                      r"too (long|short)|try again|denied|unauthori[sz]ed|forbidden|success(fully)?|saved|created|generated)\b", re.I)
_DOC_CLASSES = {"Document Generator"}


def collect_evidence(execution_run: Optional[Dict[str, Any]], action_plan: Optional[Dict[str, Any]],
                     test_cases: Sequence[Dict[str, Any]], artifact_dir: Path,
                     requirements: Sequence[Dict[str, Any]] = (), execution_error: Optional[str] = None,
                     no_plan_reason: Optional[str] = None) -> Dict[str, Any]:
    """execution_error: why the executor produced no results (e.g. the browser crashed) — the reason every planned
    test carries instead of a bare "not executed". no_plan_reason: why there is no action plan at all."""
    t0 = time.monotonic()
    results = (execution_run or {}).get("results") or []
    plans = {p["plan_id"]: p for p in (action_plan or {}).get("plans") or []}
    tests = {t.get("id"): t for t in test_cases if t.get("id")}
    statements = {c["code"]: c.get("statement") for r in requirements for c in (r.get("criteria") or []) if c.get("code")}

    bundles = [_bundle(r, plans.get(r["plan_id"]) or {}, tests.get(r.get("test_id")) or {}, artifact_dir, statements)
               for r in results]
    # tests that never got a plan (unmappable / blocked) are evidence of non-execution too
    executed = {b["test_id"] for b in bundles if b["test_id"]}
    for p in plans.values():
        if p.get("kind") == "brd_test" and p.get("test_id") not in executed:
            reason = p.get("reason") or execution_error or (action_plan or {}).get("reason")                 or "the test was planned but the browser run produced no result for it"
            bundles.append(_not_executed(p, tests.get(p.get("test_id")) or {}, statements, reason))
    if not plans:
        why = ((action_plan or {}).get("reason") if action_plan else None) or no_plan_reason             or "no browser actions were generated for this test"
        for t in test_cases:
            if t.get("id") in executed:
                continue
            bundles.append(_not_executed({"plan_id": None, "test_id": t.get("id"), "test_code": t.get("code"),
                                          "criterion_code": t.get("criterion_code"), "requirement_ref": t.get("requirement_ref"),
                                          "title": t.get("title"), "kind": "brd_test", "scored": True},
                                         t, statements, why))

    by_criterion: Dict[str, List[str]] = {}
    for b in bundles:
        if b["criterion_code"] and b["scored"]:
            by_criterion.setdefault(b["criterion_code"], []).append(b["bundle_id"])
    counts = {
        "bundles": len(bundles),
        "executed": sum(1 for b in bundles if b["execution_status"] in ("completed", "failed")),
        "with_output": sum(1 for b in bundles if b["completeness"]["output_captured"]),
        "documents": sum(len(b["observed"]["documents"]) for b in bundles),
        "not_executed": sum(1 for b in bundles if b["execution_status"] not in ("completed", "failed")),
        "criteria_covered": len(by_criterion),
    }
    collection = {"status": "collected" if bundles else "empty", "counts": counts, "by_criterion": by_criterion,
                  "bundles": bundles, "duration_s": round(time.monotonic() - t0, 2)}
    if artifact_dir.exists():
        manifest = {"bundles": [{k: b[k] for k in ("bundle_id", "plan_id", "test_code", "criterion_code", "artifacts", "digest")}
                                for b in bundles]}
        (artifact_dir / "evidence_manifest.json").write_text(json.dumps(manifest, indent=1), encoding="utf-8")
        collection["manifest"] = "evidence_manifest.json"
    return collection


# ── one bundle ──────────────────────────────────────────────────────────────

def _bundle(r: Dict[str, Any], plan: Dict[str, Any], test: Dict[str, Any], root: Path, statements: Dict) -> Dict[str, Any]:
    caps, obs = r.get("captures") or {}, r.get("observations") or {}
    artifacts: List[Dict[str, Any]] = []

    def add(kind: str, rel: Optional[str]) -> Optional[Path]:
        if not rel:
            return None
        path = root / rel
        if not path.is_file():
            return None
        data = path.read_bytes()
        artifacts.append({"type": kind, "path": rel, "size": len(data), "sha256": hashlib.sha256(data).hexdigest()})
        return path

    add("screenshot", caps.get("screenshot"))
    after_path = add("page_text", caps.get("dom_text"))
    before_path = add("page_text_before", caps.get("dom_before"))
    for s in r.get("steps") or []:
        add("failure_screenshot", s.get("screenshot"))

    after = _read_text(after_path)
    before = _read_text(before_path)
    appeared, change = _appeared(before, after) if after is not None and before is not None else (None, None)
    messages = _messages(appeared if appeared is not None else (after or ""))

    documents = []
    for d in caps.get("downloads") or []:
        path = add("download", d.get("file"))
        if path is None:
            documents.append({**d, "readable": False, "error": "file missing on disk"})
            continue
        doc = read_document(path)
        if doc.get("text"):
            text_rel = f"{d['file']}.txt"
            (root / text_rel).write_text(doc["text"][:TEXT_LIMIT], encoding="utf-8")
            add("document_text", text_rel)
            doc["text_file"] = text_rel
        doc["text_excerpt"] = (doc.pop("text", "") or "")[:EXCERPT]
        documents.append({**{k: d.get(k) for k in ("file", "name", "size", "sha256", "kind")}, **doc})

    views = []
    for v in caps.get("views") or []:
        vpath = add("result_view_text", v.get("text_file"))
        add("result_view_screenshot", v.get("screenshot"))
        vtext = _read_text(vpath) or ""
        views.append({"name": v.get("name"), "text_file": v.get("text_file"), "screenshot": v.get("screenshot"),
                      "text_excerpt": vtext[:EXCERPT], "characters": len(vtext)})

    resp = caps.get("response")
    if resp:
        add("http_response", resp.get("file"))

    observed = {
        "ui_output": (appeared or "")[:EXCERPT] if appeared else None,
        "page_change": change,  # added | reordered | none | None (no before-snapshot)
        "page_text_excerpt": (after or caps.get("dom_excerpt") or "")[:EXCERPT] or None,
        "messages": messages,
        "reply": caps.get("reply"),
        "documents": documents,
        "views": views,  # tabs of the generated result (BRD / TSD / diagrams …), each opened and captured
        "http_response": {k: resp.get(k) for k in ("status", "ms", "content_type", "excerpt")} if resp else None,
        "final_url": r.get("final_url"),
    }
    cls = plan.get("class") or r.get("class")
    bundle = {
        **_identity(r, plan, test, statements),
        "execution_status": r["status"],
        "failed_step": r.get("failed_step"),
        "error": r.get("error"),
        "duration_ms": r.get("duration_ms"),
        "inputs": _inputs(plan),
        "observed": observed,
        "expectations": r.get("expectations") or [],
        "artifacts": artifacts,
        "observations": _observations(obs),
        "completeness": _completeness(r, plan, cls, observed, artifacts),
    }
    bundle["digest"] = hashlib.sha256(json.dumps([a["sha256"] for a in artifacts] + [bundle["bundle_id"]]).encode()).hexdigest()
    return bundle


def _identity(r: Dict, plan: Dict, test: Dict, statements: Dict) -> Dict[str, Any]:
    crit = r.get("criterion_code") or plan.get("criterion_code") or test.get("criterion_code")
    return {
        "bundle_id": f"EV-{r.get('test_code') or plan.get('test_code') or r.get('plan_id') or 'X'}",
        "plan_id": r.get("plan_id"), "test_id": r.get("test_id") or test.get("id"),
        "test_code": r.get("test_code") or test.get("code"), "criterion_code": crit,
        "criterion_statement": statements.get(crit), "requirement_ref": r.get("requirement_ref") or test.get("requirement_ref"),
        "title": r.get("title") or test.get("title"), "variant": r.get("variant") or test.get("variant_type"),
        "class": r.get("class") or plan.get("class"),
        "kind": r.get("kind") or plan.get("kind") or "brd_test", "scored": bool(r.get("scored", plan.get("scored", True))),
        "test_input": test.get("input"), "expected_behavior": test.get("expected_behavior"),
    }


def _not_executed(plan: Dict, test: Dict, statements: Dict, reason: str) -> Dict[str, Any]:
    ident = _identity({"plan_id": plan.get("plan_id"), "test_id": plan.get("test_id"), "test_code": plan.get("test_code"),
                       "criterion_code": plan.get("criterion_code"), "requirement_ref": plan.get("requirement_ref"),
                       "title": plan.get("title"), "kind": plan.get("kind"), "scored": plan.get("scored", True)},
                      plan, test, statements)
    return {**ident, "execution_status": "not_executed", "failed_step": None, "error": reason, "duration_ms": 0,
            "inputs": [], "observed": {"ui_output": None, "page_change": None, "page_text_excerpt": None, "messages": [], "reply": None,
                                       "documents": [], "views": [], "http_response": None, "final_url": None},
            "expectations": [{"expectation": test.get("expected_behavior"), "pass_criteria": test.get("pass_criteria") or []}],
            "artifacts": [], "observations": {}, "digest": None,
            "completeness": {"output_captured": False, "expected_output": None, "missing": [f"not executed: {reason}"]}}


def _inputs(plan: Dict) -> List[Dict[str, Any]]:
    out = []
    for s in plan.get("steps") or []:
        if s["action"] in ("fill", "select", "check", "upload", "send_message"):
            t = s.get("target") or {}
            out.append({"step": s.get("n"), "action": s["action"], "field": t.get("label") or t.get("name") or t.get("text"),
                        "value": (s.get("value") or ("checked" if s["action"] == "check" else None)),
                        "synthetic": bool(s.get("synthetic"))})
        elif s["action"] == "http":
            out.append({"step": s.get("n"), "action": "http", "field": f"{s.get('method')} {s.get('path')}",
                        "value": json.dumps(s.get("body"), ensure_ascii=False)[:4000], "synthetic": False})
    return out


def _observations(obs: Dict) -> Dict[str, Any]:
    net = obs.get("network") or []
    slow = max(net, key=lambda n: n["ms"], default=None)
    return {"api_calls": len(net), "api_failures": [n for n in net if n["status"] >= 400 or n["status"] == 0][:10],
            "slowest_call": slow, "console_errors": (obs.get("console_errors") or [])[:10],
            "page_errors": (obs.get("page_errors") or [])[:10]}


def _completeness(r: Dict, plan: Dict, cls: Optional[str], observed: Dict, artifacts: List[Dict]) -> Dict[str, Any]:
    steps = plan.get("steps") or []
    wants_doc = any(s["action"] == "expect_download" for s in steps) or (cls in _DOC_CLASSES and r.get("variant") == "positive")
    wants_reply = any(s["action"] == "send_message" for s in steps)
    wants_http = any(s["action"] == "http" for s in steps)
    expected = "document" if wants_doc else "reply" if wants_reply else "http_response" if wants_http else "ui_output"
    missing = []
    if not any(a["type"] == "screenshot" for a in artifacts) and not wants_http:
        missing.append("no final screenshot")
    if wants_doc and not observed["documents"]:
        missing.append("expected a downloaded document but none was produced")
    if wants_doc and observed["documents"] and not any(d.get("readable") for d in observed["documents"]):
        missing.append("the downloaded document could not be read")
    if wants_reply and not observed["reply"]:
        missing.append("no chat reply appeared")
    if wants_http and not observed["http_response"]:
        missing.append("no HTTP response")
    if expected == "ui_output" and not observed["ui_output"] and not observed["messages"]:
        missing.append("nothing new appeared on the page after the actions")
    captured = {"document": bool(observed["documents"]) and any(d.get("readable") for d in observed["documents"]),
                "reply": bool(observed["reply"]), "http_response": bool(observed["http_response"]),
                "ui_output": bool(observed["ui_output"] or observed["messages"])}[expected]
    return {"output_captured": captured, "expected_output": expected, "missing": missing}


# ── text helpers ────────────────────────────────────────────────────────────

def _read_text(path: Optional[Path]) -> Optional[str]:
    if path is None:
        return None
    try:
        return path.read_text(encoding="utf-8")
    except Exception:
        return None


def _appeared(before: str, after: str) -> Tuple[str, str]:
    """What changed on the page: lines added or moved, in page order, and the kind of change.

    Order-aware (difflib), so a re-sorted list counts as a change ("reordered") even though no new line appeared.
    """
    a = [l.strip() for l in before.splitlines() if l.strip()]
    b = [l.strip() for l in after.splitlines() if l.strip()]
    changed: List[str] = []
    for op, _, _, j1, j2 in SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if op in ("replace", "insert"):
            changed += b[j1:j2]
    if not changed:
        return "", "none"
    kind = "reordered" if sorted(a) == sorted(b) else "added"
    return "\n".join(changed)[:TEXT_LIMIT], kind


def _messages(text: str) -> List[str]:
    out = []
    for line in text.splitlines():
        line = line.strip()
        if 3 <= len(line) <= 300 and _MESSAGE.search(line) and line not in out:
            out.append(line)
    return out[:12]


# ── documents ───────────────────────────────────────────────────────────────

def read_document(path: Path) -> Dict[str, Any]:
    """Open a downloaded file and describe it: text, headings, tables, pages, word count. Never raises."""
    data = path.read_bytes()
    ext = path.suffix.lower()
    try:
        if data[:8] == b"\x89PNG\r\n\x1a\n" or data[:3] == b"\xff\xd8\xff" or data[:6] in (b"GIF87a", b"GIF89a") \
                or (data[:4] == b"RIFF" and data[8:12] == b"WEBP"):
            return _image(data, ext)
        if data[:4] == b"%PDF" or ext == ".pdf":
            doc = _pdf(path)
        elif data[:2] == b"PK" and ext == ".docx":
            doc = _docx(path)
        elif data[:2] == b"PK" and ext == ".xlsx":
            doc = _xlsx(path)
        elif data[:2] == b"PK" and ext == ".pptx":
            doc = _pptx(path)
        elif data[:2] == b"PK":
            with zipfile.ZipFile(path) as z:
                names = z.namelist()
            doc = {"format": "zip", "entries": names[:50], "text": "\n".join(names)}
        else:
            text = data.decode("utf-8", errors="strict")
            doc = _plain(text, ext)
    except Exception as exc:
        return {"readable": False, "format": ext.lstrip(".") or "unknown", "error": f"{type(exc).__name__}: {exc}"[:300]}
    text = doc.get("text") or ""
    doc.update(readable=True, words=len(re.findall(r"\w+", text)), characters=len(text),
               empty=not text.strip())
    doc.setdefault("headings", [])
    doc["headings"] = doc["headings"][:60]
    return doc


def _image(data: bytes, ext: str) -> Dict[str, Any]:
    """A downloaded diagram / picture: readable, no text; PNG dimensions when available."""
    fmt = "png" if data[:4] == b"\x89PNG" else "jpeg" if data[:2] == b"\xff\xd8" else "gif" if data[:3] == b"GIF" else "webp"
    out: Dict[str, Any] = {"readable": True, "format": fmt, "kind": "image", "words": 0, "characters": 0, "empty": len(data) == 0,
                           "headings": [], "bytes": len(data)}
    if fmt == "png" and len(data) >= 24:
        out["width"], out["height"] = int.from_bytes(data[16:20], "big"), int.from_bytes(data[20:24], "big")
    return out


def _docx(path: Path) -> Dict[str, Any]:
    import docx

    d = docx.Document(str(path))
    paras, headings = [], []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        paras.append(t)
        style = (p.style.name if p.style is not None else "") or ""
        if style.lower().startswith(("heading", "title")):
            headings.append(t)
    table_text = ["\t".join(c.text.strip() for c in row.cells) for tbl in d.tables for row in tbl.rows]
    with zipfile.ZipFile(path) as z:  # embedded pictures (e.g. rendered diagrams) live in word/media/
        images = sum(1 for n in z.namelist() if n.startswith("word/media/"))
    return {"format": "docx", "text": "\n".join(paras + table_text), "headings": headings, "tables": len(d.tables),
            "paragraphs": len(paras), "images": images}


def _pdf(path: Path) -> Dict[str, Any]:
    import fitz

    with fitz.open(str(path)) as d:
        text = "\n".join(page.get_text() for page in d)
        toc = [t[1] for t in d.get_toc()]
        pages = d.page_count
    return {"format": "pdf", "text": text, "headings": toc, "pages": pages}


def _xlsx(path: Path) -> Dict[str, Any]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    lines, sheets = [], []
    for ws in wb.worksheets:
        rows = 0
        for row in ws.iter_rows(values_only=True):
            rows += 1
            if len(lines) < 5000:
                lines.append("\t".join("" if v is None else str(v) for v in row))
        sheets.append({"name": ws.title, "rows": rows})
    wb.close()
    return {"format": "xlsx", "text": "\n".join(lines), "headings": [s["name"] for s in sheets], "sheets": sheets}


def _pptx(path: Path) -> Dict[str, Any]:
    """No python-pptx needed: slide text is in ppt/slides/slideN.xml <a:t> runs."""
    with zipfile.ZipFile(path) as z:
        slides = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                        key=lambda n: int(re.findall(r"\d+", n)[-1]))
        texts, titles = [], []
        for n in slides:
            runs = re.findall(r"<a:t>([^<]*)</a:t>", z.read(n).decode("utf-8", errors="replace"))
            texts.append(" ".join(runs))
            if runs:
                titles.append(runs[0])
    return {"format": "pptx", "text": "\n".join(texts), "headings": titles, "slides": len(slides)}


def _plain(text: str, ext: str) -> Dict[str, Any]:
    fmt = ext.lstrip(".") or "text"
    headings: List[str] = []
    if fmt in ("md", "markdown", "txt", "text"):
        headings = [l.lstrip("#").strip() for l in text.splitlines() if re.match(r"^#{1,6}\s+\S", l)]
    if fmt == "json":
        json.loads(text)  # unreadable JSON is recorded as such by the caller
    if fmt in ("html", "htm"):
        headings = [re.sub(r"<[^>]+>", "", h).strip() for h in re.findall(r"<h[1-3][^>]*>(.*?)</h[1-3]>", text, re.I | re.S)]
        text = re.sub(r"<[^>]+>", " ", text)
    return {"format": fmt, "text": text, "headings": headings}
