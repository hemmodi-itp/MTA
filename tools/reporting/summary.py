"""Phase 11: executive summary metrics, traceability CSV, defect list and quality gates. Pure Python.
(Risk areas, trend, coverage dashboard, exports and the 29119-3 summary: tools/reporting/insights.py.)

Flags in project.yaml (absent = OFF):
    reporting:     {enhanced: true}                       # executive summary / defects / CSV in the report
    quality_gates:                                        # any subset; a failed gate -> exit code 1
      p0_pass_rate: 100            # minimum %, P0 tests
      p1_pass_rate: 95             # minimum %, P1 tests
      requirement_coverage: 90     # minimum %, automatable ACs with a test
      max_critical_a11y: 0         # maximum critical/serious axe violations (from the attached findings)
      max_security_high: 0         # maximum failing passive-security findings
      max_flaky_rate: 2            # maximum % of tests that were flaky
      perf_budget_regression_pct: 20   # maximum % a performance metric got worse than the baseline

CLI: python -m tools.reporting.summary <report.json>   (exit 0 = gates passed or none configured, 1 = failed)
"""

import csv
import io
import json
import os
import sys
from typing import Any, Dict, List, Optional

PRIORITIES = ("P0", "P1", "P2", "P3")
_MIN_GATES = ("p0_pass_rate", "p1_pass_rate", "requirement_coverage")
_MAX_GATES = ("max_critical_a11y", "max_security_high", "max_flaky_rate", "perf_budget_regression_pct")
_GATE_KEYS = _MIN_GATES + _MAX_GATES


def load_project_yaml(project_dir: str) -> dict:
    import yaml
    try:
        with open(os.path.join(project_dir, "project.yaml"), encoding="utf-8") as fh:
            return yaml.safe_load(fh) or {}
    except (OSError, ValueError, yaml.YAMLError):
        return {}


def enhanced_enabled(raw: dict) -> bool:
    return bool((raw.get("reporting") or {}).get("enhanced"))


def gates_config(raw: dict) -> Dict[str, float]:
    block = raw.get("quality_gates")
    if not isinstance(block, dict):
        return {}
    return {k: float(block[k]) for k in _GATE_KEYS if isinstance(block.get(k), (int, float))}


def _is_pass(r: dict) -> bool:
    return str(r.get("pw_status") or r.get("status") or "").lower() in ("passed", "expected", "success")


def _is_skipped(r: dict) -> bool:
    return str(r.get("pw_status") or r.get("status") or "").lower() == "skipped"


def _priority_by_tc(plans: dict) -> Dict[str, str]:
    out = {}
    for plan in ((plans or {}).get("plans") or {}).values():
        for t in (plan.get("tests") or []) if isinstance(plan, dict) else []:
            if isinstance(t, dict) and t.get("tc_id") and t.get("priority"):
                out[t["tc_id"]] = t["priority"]
    return out


def pass_rate_by_priority(results: List[dict], tc_priority: Dict[str, str]) -> Dict[str, Optional[dict]]:
    """{P0: {passed, executed, rate} | None}; skipped tests are not executed. Unprioritised tests are ignored."""
    acc = {p: [0, 0] for p in PRIORITIES}
    for r in results:
        p = tc_priority.get(r.get("tc_id") or "")
        if p not in acc or _is_skipped(r):
            continue
        acc[p][1] += 1
        acc[p][0] += 1 if _is_pass(r) else 0
    return {p: ({"passed": a, "executed": n, "rate": round(100.0 * a / n, 1)} if n else None)
            for p, (a, n) in acc.items()}


def _plan_tests(plans: dict) -> List[dict]:
    return [t for plan in ((plans or {}).get("plans") or {}).values() if isinstance(plan, dict)
            for t in plan.get("tests") or [] if isinstance(t, dict)]


def potential_defects(test_cases: Optional[dict]) -> List[dict]:
    """Phase 5.2: BRD-vs-observed conflicts from test design — potential defects awaiting a human decision."""
    return [{"ac_id": c.get("ac_id"), "where": c.get("form_id"), "field": c.get("field"), "probe": c.get("probe"),
             "brd": c.get("brd"), "observed": c.get("observed")}
            for c in (test_cases or {}).get("conflicts") or []]


def compute_summary(results: List[dict], trace: Optional[dict], plans: dict, suite: dict,
                    triage: Optional[dict], agent_quality: Optional[dict],
                    test_cases: Optional[dict] = None) -> Dict[str, Any]:
    ts = (trace or {}).get("summary") or {}
    weak = sum(len(s.get("weak_oracle") or []) for s in (suite or {}).get("specs") or [])
    failed = [r for r in results if not _is_pass(r) and not _is_skipped(r)
              and str(r.get("pw_status") or "").lower() != "flaky"]
    tag_fail = lambda tag: sum(1 for r in failed if tag in (r.get("tags") or []))  # noqa: E731
    dims = ((agent_quality or {}).get("dimensions") or {})
    from tools.reporting.insights import nfr_counts
    nfr = nfr_counts(results)
    has_findings = any(str(a.get("name")) in ("a11y-findings.json", "security-findings.json")
                       for r in results for a in r.get("attachments") or [])
    flaky = sum(1 for r in results if str(r.get("pw_status") or r.get("status") or "").lower() == "flaky")
    return {
        "requirement_coverage": ts.get("coverage_pct"),
        "acs_covered": ts.get("acs_covered"), "acs_automatable": ts.get("acs_automatable"),
        "pass_rate_by_priority": pass_rate_by_priority(results, _priority_by_tc(plans)),
        "product_defects": (triage["counts"].get("product_defect", 0) if triage else None),
        # counted from the findings the NFR specs attach; failed @a11y/@security tests only as a fallback
        "a11y_findings": nfr["critical_a11y"] if has_findings else tag_fail("@a11y"),
        "security_findings": nfr["high_security"] if has_findings else tag_fail("@security"),
        "critical_a11y": nfr["critical_a11y"], "high_security": nfr["high_security"],
        "perf_regression_pct": nfr["perf_regression_pct"],
        "flaky_rate": round(100.0 * flaky / len(results), 1) if results else None,
        "weak_oracle": weak,
        # Phase 5.2: expected values only an LLM inferred (reported, not BRD-backed) and BRD-vs-observed conflicts
        "low_confidence_oracles": sum(1 for t in _plan_tests(plans) if t.get("oracle_confidence") == "low"),
        "potential_defects": len(potential_defects(test_cases)),
        "agent_dimensions": {d: v.get("outcome") for d, v in dims.items()} if dims else None,
    }


def evaluate_gates(summary: dict, gates: Dict[str, float]) -> Dict[str, Any]:
    """{verdict: PASS|FAIL|NO_GATES, gates: [{gate, threshold, actual, ok, kind}]}. Missing data fails (no silent
    pass), except a P0/P1 rate with zero executed tests of that priority, and a perf regression with no perf run."""
    rows = []
    for gate, thr in gates.items():
        if gate in ("p0_pass_rate", "p1_pass_rate"):
            p = summary["pass_rate_by_priority"].get(gate[:2].upper())
            actual = p["rate"] if p else None
            ok = True if p is None else actual >= thr
        elif gate == "requirement_coverage":
            actual = summary.get("requirement_coverage")
            ok = actual is not None and actual >= thr
        elif gate == "max_critical_a11y":
            actual = summary.get("critical_a11y")
            ok = actual is not None and actual <= thr
        elif gate == "max_security_high":
            actual = summary.get("high_security")
            ok = actual is not None and actual <= thr
        elif gate == "max_flaky_rate":
            actual = summary.get("flaky_rate")
            ok = actual is not None and actual <= thr
        else:   # perf_budget_regression_pct
            actual = summary.get("perf_regression_pct")
            ok = actual is None or actual <= thr
        rows.append({"gate": gate, "threshold": thr, "actual": actual, "ok": ok,
                     "kind": "min" if gate in _MIN_GATES else "max"})
    verdict = "NO_GATES" if not rows else ("PASS" if all(r["ok"] for r in rows) else "FAIL")
    return {"verdict": verdict, "gates": rows}


def gate_exit_code(gate_result: Optional[dict]) -> int:
    return 1 if gate_result and gate_result.get("verdict") == "FAIL" else 0


def traceability_csv(trace: Optional[dict]) -> str:
    """AC × tests × latest result; an AC with no test gets one row saying why."""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["ac_id", "priority", "criterion", "ac_status", "tc_id", "test_title", "latest_result", "not_tested_reason"])
    for a in (trace or {}).get("acs") or []:
        if a.get("children"):
            continue
        reason = a.get("uncovered_reason") or {}
        base = [a.get("id"), a.get("priority"), a.get("statement"), a.get("status")]
        tests = a.get("tests") or []
        if not tests:
            w.writerow(base + ["", "", "not_run", reason.get("code", "")])
        for t in tests:
            w.writerow(base + [t.get("tc_id"), t.get("title"), t.get("status"), ""])
    return buf.getvalue()


def traceability_xlsx(trace: Optional[dict], path: str) -> Optional[str]:
    """The traceability matrix as .xlsx (same rows as the CSV) with a header filter; None when openpyxl is missing."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return None
    rows = list(csv.reader(io.StringIO(traceability_csv(trace))))
    wb = Workbook()
    ws = wb.active
    ws.title = "Traceability"
    for row in rows:
        ws.append(row)
    fill = {"failed": "FECACA", "passed": "BBF7D0", "flaky": "FDE68A", "not_run": "E5E7EB"}
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in ws.iter_rows(min_row=2):
        color = fill.get(str(r[6].value or ""))
        if color:
            r[6].fill = PatternFill("solid", fgColor=color)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    for col, width in zip("ABCDEFGH", (12, 9, 60, 12, 16, 50, 12, 22)):
        ws.column_dimensions[col].width = width
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
    return path


def not_tested(trace: Optional[dict]) -> List[dict]:
    return [{"id": u.get("id"), "priority": u.get("priority"), "code": u.get("code") or "unknown",
             "detail": u.get("detail") or ""} for u in (trace or {}).get("uncovered") or []]


def _step_text(s: dict) -> str:
    parts = [str(s.get("action", "")), str(s.get("locator_key") or s.get("target") or ""), str(s.get("value") or "")]
    return " ".join(p for p in parts if p)


def _plan_test(plans: dict, tc_id: str, title: str = "") -> Optional[dict]:
    """The plan test of a result: by tc_id (traced projects), else by its rendered title (untraced)."""
    from tools.test_creation.script_generator.spec_renderer import render_test_title
    tests = [t for plan in ((plans or {}).get("plans") or plans or {}).values() if isinstance(plan, dict)
             for t in plan.get("tests") or [] if isinstance(t, dict)]
    if tc_id:
        hit = next((t for t in tests if t.get("tc_id") == tc_id), None)
        if hit:
            return hit
    if title:
        for t in tests:
            try:
                if t.get("title") == title or (t.get("name") and render_test_title(t) == title):
                    return t
            except KeyError:
                continue
    return None


def _first_line(errors: list) -> str:
    for e in errors or []:
        line = str(e or "").strip().splitlines()
        if line:
            return line[0][:300]
    return ""


def _expected_source(plan_test: dict) -> str:
    """Where the expected value came from (plan 5.2 precedence) — BRD unless the oracle was only LLM-inferred."""
    if plan_test.get("oracle_confidence") == "low":
        return "llm_inferred (low confidence)"
    return "brd" if plan_test.get("ac_ids") else "observed / generated"


def build_defects(results: List[dict], triage: Optional[dict], plans: dict, run: Optional[str] = None,
                  environment: str = "") -> List[dict]:
    """One entry per triaged product_defect: title, linked ACs, steps (from the plan), expected vs actual, evidence."""
    if not triage:
        return []
    from tools.triage.triage import PRODUCT_DEFECT, _spec_key
    # triage.json already holds spec keys ("M02/a.spec.ts"); a result holds Playwright's path — compare keys, never
    # basenames (two modules may have a spec with the same file name)
    bad = {(str(t["spec_file"]).replace("\\", "/"), t["test_name"]) for t in triage["tests"]
           if t["verdict"] == PRODUCT_DEFECT}
    out = []
    for r in results:
        if (_spec_key(r.get("spec_file") or ""), r.get("test_title")) not in bad:
            continue
        pt = _plan_test(plans, r.get("tc_id") or "", r.get("test_title") or "") or {}
        steps = [_step_text(s) for s in pt.get("steps") or [] if isinstance(s, dict)]
        expected = "; ".join(f"{s.get('message') or _step_text(s)}" for s in pt.get("steps") or []
                             if isinstance(s, dict) and str(s.get("action", "")).startswith("verify"))
        atts = r.get("attachments") or []
        link = lambda kind: next((a.get("path") for a in atts if kind in str(a.get("name", "")).lower()  # noqa: E731
                                  or kind in str(a.get("content_type", "")).lower()), None)
        from tools.reporting.insights import evidence, severity_of
        ev = evidence(r, run=run, environment=environment)
        out.append({"title": r.get("test_title"), "tc_id": r.get("tc_id"), "severity": severity_of(r, pt),
                    "priority": pt.get("priority"),
                    "ac_ids": r.get("ac_ids") or pt.get("ac_ids") or [], "steps": steps,
                    "expected": expected, "expected_source": _expected_source(pt), "actual": _first_line(r.get("errors")),
                    "screenshot": link("screenshot") or link("image"), "trace": link("trace"), "evidence": ev})
    return out


def main(argv: Optional[List[str]] = None) -> int:
    argv = sys.argv[1:] if argv is None else argv
    if not argv:
        print(__doc__)
        return 2
    with open(argv[0], encoding="utf-8") as fh:
        gates = (json.load(fh).get("quality") or {}).get("gates")
    if gates:
        for g in gates["gates"]:
            op = "<=" if g.get("kind") == "max" else ">="
            print(f"{'PASS' if g['ok'] else 'FAIL'} {g['gate']}: {g['actual']} (need {op} {g['threshold']})")
    return gate_exit_code(gates)


if __name__ == "__main__":
    sys.exit(main())
