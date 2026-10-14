"""Phase 11 (complete): severity, risk areas, history/trend, coverage dashboard, NFR counts from attached findings,
evidence, Jira/ADO exports, 29119-3 summary, all quality gates, XLSX matrix, HTML sections, gate exit code."""
import json

import pytest

from tools.reporting import insights, summary as qs


def R(title, status="unexpected", spec="test_creation/test_scripts/M02/a.spec.ts", **kw):
    return {"test_title": title, "pw_status": status, "spec_file": spec, **kw}


PLANS = {"plans": {"S1": {"tests": [{"tc_id": "TC-1", "priority": "P0", "steps": [], "ac_ids": ["AC-1.1"]},
                                    {"tc_id": "TC-2", "priority": "P1", "steps": []},
                                    {"tc_id": "TC-3", "priority": "P3", "oracle_confidence": "low", "steps": []}]}}}


def test_severity_from_priority_and_nfr_tags():
    idx = insights._plan_index(PLANS)
    assert insights.severity_of(R("x"), idx["TC-1"]) == "critical"
    assert insights.severity_of(R("x"), idx["TC-3"]) == "low"
    assert insights.severity_of(R("a11y: /x", tags=["@a11y"]), None) == "high"
    assert insights.severity_of(R("x"), None) == "medium"
    assert insights.defects_by_severity([{"severity": "high"}, {"severity": "high"}, {}]) == \
        {"critical": 0, "high": 2, "medium": 1, "low": 0}


def test_risk_areas_rank_failures_untested_and_flaky():
    results = [R("a", tc_id="TC-1"), R("b", tc_id="TC-2"), R("c", "flaky", spec="x/test_scripts/M05/z.spec.ts"),
               R("d", "expected", spec="x/test_scripts/M09/z.spec.ts")]
    trace = {"uncovered": [{"id": "AC-3.2", "priority": "P0"}, {"id": "AC-3.4", "priority": "P3"}]}
    risks = insights.risk_areas(results, trace, PLANS)
    assert [r["area"] for r in risks] == ["M02", "requirement AC-3", "M05"]
    assert risks[0]["score"] == 8 + 4 and "2 failing test(s)" in risks[0]["why"]


def test_history_is_idempotent_per_run_and_trend_has_deltas(tmp_path):
    res = [R("a", "expected"), R("b", "unexpected")]
    insights.record_history(str(tmp_path), {"requirement_coverage": 50, "product_defects": 1}, None, res, run="r1")
    insights.record_history(str(tmp_path), {"requirement_coverage": 50, "product_defects": 1}, None, res, run="r1")
    rows = insights.record_history(str(tmp_path), {"requirement_coverage": 60, "product_defects": 0},
                                   {"verdict": "PASS"}, [R("a", "expected")], run="r2")
    assert [r["run"] for r in rows] == ["r1", "r2"]
    t = insights.trend(rows)
    assert t["runs"] == rows[:1] and t["delta"] == {"pass_rate": 50.0, "requirement_coverage": 10, "product_defects": -1,
                                                    "flaky": 0}


def test_coverage_dashboard():
    spec = 'await action.click("save"); await page.goto("http://h:1/products/7"); await action.navigate("/cart"); ' \
           'verifyApiState("GET", "/api/orders/2")'
    cov = insights.coverage_dashboard(
        [R("a", "expected", tc_id="TC-1", browser="chromium"), R("b", tags=["@a11y"], browser="webkit")],
        {"summary": {"acs_covered": 3, "acs_automatable": 4, "coverage_pct": 75.0}}, PLANS,
        {"test_cases": [{"technique": "bva"}, {"technique": "role_matrix", "roles": ["guest", "user"]},
                        {"technique": "state_transition"}], "cut": [{"technique": "bva"}]},
        {"save": {"type": "css"}, "other": {"type": "css"}},
        {"states": [{"url_template": "/products/{id}"}, {"url_template": "/cart"}, {"url_template": "/admin"}],
         "transitions": [1, 2]},
        {"endpoints": [{"path_template": "/api/orders/{id}", "kind": "fetch"}, {"path_template": "/api/x", "kind": "xhr"},
                       {"path_template": "/cart", "kind": "page"}]},
        {"a.spec.ts": spec})
    assert cov["element"] == {"used": 1, "total": 2, "pct": 50.0}
    assert cov["page_state"]["tested"] == 2 and cov["page_state"]["untested"] == ["/admin"]
    assert cov["api"] == {"referenced": 1, "total": 2, "pct": 50.0}
    assert cov["technique"]["bva"] == {"selected": 1, "cut": 1}
    assert cov["role_matrix"]["roles"] == ["guest", "user"] and cov["transition"] == {"cases": 1, "known": 2}
    assert cov["browser"] == {"chromium": 1, "webkit": 1}
    assert cov["test_type"] == {"P0".lower() if False else "functional": 1, "nfr:a11y": 1}
    assert "not measured" in cov["code"]


def test_nfr_counts_read_the_attached_findings(tmp_path):
    a11y = tmp_path / "a.json"
    a11y.write_text(json.dumps([{"rule": "axe:label", "message": "labels (critical, 1 node(s))", "level": "fail"},
                                {"rule": "axe:region", "message": "region (moderate, 2 node(s))", "level": "warn"},
                                {"rule": "keyboard-trap", "message": "x", "level": "fail"}]))
    sec = tmp_path / "s.json"
    sec.write_text(json.dumps([{"rule": "csp", "message": "m"}, {"rule": "target-blank", "message": "m", "level": "warn"}]))
    res = [R("a11y", attachments=[{"name": "a11y-findings.json", "path": str(a11y)}]),
           R("sec", attachments=[{"name": "security-findings.json", "path": str(sec)}]),
           R("perf @perf", annotations=[{"type": "perf-warning", "description": "regression:ttfb_ms: ttfb_ms 80 is 60% worse than the baseline 50"}],
             errors=["load_ms 900 is 120% worse than the baseline 400"], tags=["@perf"])]
    assert insights.nfr_counts(res) == {"critical_a11y": 1, "high_security": 1, "perf_regression_pct": 120}
    s = qs.compute_summary(res, None, {}, {}, None, None)
    assert s["a11y_findings"] == 1 and s["security_findings"] == 1 and s["flaky_rate"] == 0.0


@pytest.mark.parametrize("gate,summary_part,ok", [
    ("p1_pass_rate", {"pass_rate_by_priority": {"P1": {"rate": 90.0}}}, False),
    ("p1_pass_rate", {"pass_rate_by_priority": {"P1": None}}, True),
    ("max_critical_a11y", {"critical_a11y": 0}, True),
    ("max_critical_a11y", {"critical_a11y": 2}, False),
    ("max_critical_a11y", {"critical_a11y": None}, False),
    ("max_security_high", {"high_security": 1}, False),
    ("max_flaky_rate", {"flaky_rate": 1.5}, True),
    ("max_flaky_rate", {"flaky_rate": 4.0}, False),
    ("perf_budget_regression_pct", {"perf_regression_pct": 35}, False),
    ("perf_budget_regression_pct", {"perf_regression_pct": None}, True),
])
def test_every_quality_gate(gate, summary_part, ok):
    thr = {"p1_pass_rate": 95, "max_critical_a11y": 0, "max_security_high": 0, "max_flaky_rate": 2,
           "perf_budget_regression_pct": 20}[gate]
    base = {"pass_rate_by_priority": {}, **summary_part}
    out = qs.evaluate_gates(base, {gate: thr})
    assert out["gates"][0]["ok"] is ok and out["verdict"] == ("PASS" if ok else "FAIL")
    assert qs.gate_exit_code(out) == (0 if ok else 1)


def test_gates_config_reads_all_keys():
    raw = {"quality_gates": {"p0_pass_rate": 100, "p1_pass_rate": 95, "requirement_coverage": 90, "max_critical_a11y": 0,
                             "max_security_high": 0, "max_flaky_rate": 2, "perf_budget_regression_pct": 20, "bogus": 1}}
    assert set(qs.gates_config(raw)) == set(qs._GATE_KEYS) and len(qs._GATE_KEYS) == 7


def test_evidence_and_exports(tmp_path):
    har = tmp_path / "n.har"
    har.write_text(json.dumps({"log": {"entries": [
        {"request": {"method": "POST", "url": "http://h/p"}, "response": {"status": 500}},
        {"request": {"method": "GET", "url": "http://h/ok"}, "response": {"status": 200}}]}}))
    r = R("t", browser="webkit", annotations=[{"type": "role", "description": "user"}], attachments=[
        {"name": "screenshot", "path": "s.png", "content_type": "image/png"}, {"name": "trace", "path": "t.zip"},
        {"name": "video", "path": "v.webm"}, {"name": "network.har", "path": str(har)},
        {"name": "console-errors.json", "path": "c.json"}])
    ev = insights.evidence(r, run="r1", environment="staging")
    assert ev["video"] == "v.webm" and ev["console"] == "c.json" and ev["har_excerpt"] == ["POST http://h/p -> 500"]
    assert (ev["browser"], ev["role"], ev["run"], ev["environment"]) == ("webkit", "user", "r1", "staging")
    d = {"title": "Total wrong", "severity": "critical", "steps": ["open cart"], "expected": "10", "actual": "12",
         "ac_ids": ["AC-4.1"], "evidence": ev}
    jira = insights.jira_issues([d], "SHOP")[0]["fields"]
    assert jira["project"]["key"] == "SHOP" and jira["priority"]["name"] == "Highest" and "AC-4.1" in jira["labels"]
    assert "# open cart" in jira["description"]
    ado = insights.ado_work_items([d])[0]
    assert {"op": "add", "path": "/fields/Microsoft.VSTS.Common.Severity", "value": "1 - Critical"} in ado


def test_defects_carry_severity_expected_source_and_evidence():
    triage = {"tests": [{"spec_file": "M02/a.spec.ts", "test_name": "t", "verdict": "product_defect"}]}
    res = [R("t", tc_id="TC-3", errors=["boom"], attachments=[{"name": "video", "path": "v.webm"}])]
    d = qs.build_defects(res, triage, PLANS, run="r9")[0]
    assert d["severity"] == "low" and d["expected_source"].startswith("llm_inferred") and d["evidence"]["video"] == "v.webm"


def test_test_summary_report_recommendation():
    risks = [{"area": "M02", "why": ["2 failing test(s)"], "score": 9}]
    gates = {"verdict": "FAIL", "gates": [{"gate": "p0_pass_rate", "ok": False}]}
    tsr = insights.test_summary_report("shop", {}, gates, [R("a"), R("b", "expected")],
                                       [{"severity": "high"}], [{"id": "AC-1"}], {"requirement": {"total": 4}}, risks, [])
    assert tsr["recommendation"].startswith("Do not release") and "p0_pass_rate" in tsr["recommendation"]
    assert tsr["results"]["failed"] == 1 and "M02: 2 failing test(s)" in tsr["residual_risk"]
    md = insights.summary_markdown(tsr, "shop")
    for h in ("## Scope", "## Approach", "## Deviations", "## Coverage", "## Results", "## Defects",
              "## Residual risk", "## Recommendation"):
        assert h in md
    ok = insights.test_summary_report("shop", {}, {"verdict": "PASS"}, [], [], [], {}, [], [])
    assert ok["recommendation"].startswith("Releasable: gates passed")


def test_xlsx_matrix(tmp_path):
    trace = {"acs": [{"id": "AC-1", "priority": "P0", "statement": "s", "status": "covered",
                      "tests": [{"tc_id": "TC-1", "title": "t", "status": "failed"}]},
                     {"id": "AC-2", "priority": "P1", "statement": "u", "status": "not_covered",
                      "uncovered_reason": {"code": "no_locator"}}]}
    path = qs.traceability_xlsx(trace, str(tmp_path / "m.xlsx"))
    from openpyxl import load_workbook
    ws = load_workbook(path).active
    assert ws["A1"].value == "ac_id" and ws["G2"].value == "failed" and ws["H3"].value == "no_locator"
    assert ws.auto_filter.ref and ws.freeze_panes == "A2"


def test_html_sections_link_numbers_to_evidence_and_matrix_is_interactive():
    from tools.test_execution.reports.html_report import _quality_sections, _traceability_section, _JS
    q = {"summary": {"requirement_coverage": 50.0, "product_defects": 1, "defects_by_severity": {"high": 1},
                     "pass_rate_by_priority": {}}, "gates": {"verdict": "FAIL", "gates": [
        {"gate": "max_flaky_rate", "actual": 5, "threshold": 2, "ok": False, "kind": "max"}]},
         "not_tested": [{"id": "AC-2", "priority": "P1", "code": "no_locator", "detail": ""}],
         "defects": [{"title": "T", "severity": "high", "steps": ["s"], "expected": "e", "actual": "a",
                      "evidence": {"video": "v.webm", "browser": "webkit", "har_excerpt": ["POST /p -> 500"]}}],
         "risk_areas": [{"area": "M02", "score": 4, "why": ["1 failing test(s)"]}],
         "trend": {"runs": [{"run": "r1", "pass_rate": 90}], "delta": {"pass_rate": -10}},
         "coverage": {"requirement": {"covered": 1, "total": 2, "pct": 50.0}, "code": "not measured"},
         "test_summary_report": {"scope": "s", "approach": "a", "deviations": [], "results": {}, "defects": {},
                                 "residual_risk": [], "verdict": "FAIL", "recommendation": "Do not release"},
         "xlsx_path": "../csv/m.xlsx", "exports": {"jira": "../defects/j.json"}}
    html = _quality_sections(q)
    for frag in ('href="#traceability"', 'href="#defects"', 'href="#not-tested"', 'id="risk"', 'id="trend"',
                 'id="coverage"', 'id="tsr"', "(need ≤ 2)", 'href="v.webm"', "POST /p -&gt; 500", "jira export",
                 'href="../csv/m.xlsx"', "(-10)"):
        assert frag in html, frag
    tr = _traceability_section({"acs": [{"id": "AC-1", "tests": []}], "summary": {}})
    assert 'id="trace-matrix"' in tr and "filterTable('trace-matrix'" in tr and "sortTable" in tr
    assert "function sortTable" in _JS and "function filterTable" in _JS


def test_main_gate_exit_code(capsys):
    import main
    fail = {"reporting": {"quality_gates": {"verdict": "FAIL", "gates": [
        {"gate": "requirement_coverage", "actual": 25.9, "threshold": 90, "ok": False, "kind": "min"}]}}}
    assert main.gate_exit(fail) == 1 and "requirement_coverage = 25.9 (need >= 90)" in capsys.readouterr().out
    assert main.gate_exit(fail, ignore=True) == 0
    assert main.gate_exit({"reporting": {"quality_gates": {"verdict": "PASS"}}}) == 0
    assert main.gate_exit({}) == 0
