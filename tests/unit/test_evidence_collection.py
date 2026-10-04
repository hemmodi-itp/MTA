"""Unit tests for the Evidence Collection Engine (engines/runtime/evidence.py)."""

import hashlib
import json
import zipfile

from engines.runtime.evidence import collect_evidence, read_document


def _docx(path):
    import docx

    d = docx.Document()
    d.add_heading("Business Requirements Document", 0)
    d.add_heading("1. Problem statement", 1)
    d.add_paragraph("Clinic patients cannot book appointments online.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "ID", "Requirement"
    t.cell(1, 0).text, t.cell(1, 1).text = "FR-01", "Online booking"
    d.save(str(path))


def _pdf(path):
    import fitz

    d = fitz.open()
    p = d.new_page()
    p.insert_text((72, 72), "Functional Requirements Document")
    d.set_toc([[1, "Overview", 1]])
    d.save(str(path))


def _xlsx(path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Traceability"
    wb.active.append(["REQ", "Test"])
    wb.active.append(["REQ-01", "TC-001"])
    wb.save(str(path))


def _pptx(path):
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("ppt/slides/slide1.xml", "<p:sld><a:t>Executive summary</a:t><a:t>Scope</a:t></p:sld>")
        z.writestr("ppt/slides/slide2.xml", "<p:sld><a:t>Timeline</a:t></p:sld>")


def test_read_document_formats(tmp_path):
    for name, make in (("brd.docx", _docx), ("frd.pdf", _pdf), ("trace.xlsx", _xlsx), ("deck.pptx", _pptx)):
        make(tmp_path / name)
    d = read_document(tmp_path / "brd.docx")
    assert d["readable"] and d["format"] == "docx" and d["tables"] == 1
    assert "1. Problem statement" in d["headings"] and "FR-01\tOnline booking" in d["text"]
    p = read_document(tmp_path / "frd.pdf")
    assert p["format"] == "pdf" and p["pages"] == 1 and p["headings"] == ["Overview"] and "Functional" in p["text"]
    x = read_document(tmp_path / "trace.xlsx")
    assert x["sheets"] == [{"name": "Traceability", "rows": 2}] and "REQ-01\tTC-001" in x["text"]
    s = read_document(tmp_path / "deck.pptx")
    assert s["slides"] == 2 and s["headings"] == ["Executive summary", "Timeline"]
    (tmp_path / "broken.docx").write_bytes(b"PK\x03\x04 not really a docx")
    b = read_document(tmp_path / "broken.docx")
    assert b["readable"] is False and b["error"]
    (tmp_path / "notes.md").write_text("# Title\n\n## Scope\ntext", encoding="utf-8")
    assert read_document(tmp_path / "notes.md")["headings"] == ["Title", "Scope"]


def _write(root, rel, data):
    p = root / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(data if isinstance(data, bytes) else data.encode("utf-8"))
    return rel


def _result(root, with_download=True):
    caps = {"screenshot": _write(root, "AP-TC-001/screenshot.png", b"\x89PNG fake"),
            "dom_before": _write(root, "AP-TC-001/dom_before.txt", "Create a BRD\nProject name\nGenerate documents"),
            "dom_text": _write(root, "AP-TC-001/dom.txt",
                               "Create a BRD\nProject name\nGenerate documents\nYour BRD is ready\nDocuments generated successfully"),
            "downloads": []}
    if with_download:
        (root / "AP-TC-001" / "downloads").mkdir(parents=True, exist_ok=True)
        _docx(root / "AP-TC-001" / "downloads" / "brd.docx")
        data = (root / "AP-TC-001/downloads/brd.docx").read_bytes()
        caps["downloads"] = [{"file": "AP-TC-001/downloads/brd.docx", "name": "brd.docx", "size": len(data),
                              "sha256": hashlib.sha256(data).hexdigest(), "kind": "docx"}]
    return {"plan_id": "AP-TC-001", "test_id": "t1", "test_code": "TC-001", "criterion_code": "AC-01.1", "kind": "brd_test",
            "scored": True, "class": "Document Generator", "variant": "positive", "status": "completed", "duration_ms": 9000,
            "steps": [{"n": 1, "action": "goto", "status": "passed"}], "captures": caps,
            "observations": {"network": [{"method": "POST", "path": "/api/generate", "status": 200, "ms": 41000},
                                         {"method": "GET", "path": "/api/x", "status": 500, "ms": 20}],
                             "console_errors": ["boom"], "page_errors": []},
            "expectations": [{"expectation": "A BRD is produced", "pass_criteria": []}]}


PLAN = {"status": "ready", "plans": [
    {"plan_id": "AP-TC-001", "test_id": "t1", "test_code": "TC-001", "criterion_code": "AC-01.1", "kind": "brd_test",
     "scored": True, "class": "Document Generator", "status": "ready", "steps": [
         {"n": 3, "action": "fill", "target": {"label": "Problem statement"}, "value": "Clinic booking"},
         {"n": 4, "action": "fill", "target": {"label": "Project name"}, "value": "MTA Evaluation Project", "synthetic": True},
         {"n": 7, "action": "expect_download", "target": {"text": "Download"}}]},
    {"plan_id": "AP-TC-002", "test_id": "t2", "test_code": "TC-002", "criterion_code": "AC-02.1", "kind": "brd_test",
     "scored": True, "status": "unmappable", "reason": "latency claim", "steps": []}]}
TESTS = [{"id": "t1", "code": "TC-001", "criterion_code": "AC-01.1", "input": "Clinic booking", "expected_behavior": "A BRD is produced"},
         {"id": "t2", "code": "TC-002", "criterion_code": "AC-02.1", "expected_behavior": "p95 < 2s"}]
REQS = [{"code": "REQ-01", "criteria": [{"code": "AC-01.1", "statement": "The system generates a BRD document"}]}]


def test_bundle_has_inputs_outputs_documents_and_hashes(tmp_path):
    col = collect_evidence({"results": [_result(tmp_path)]}, PLAN, TESTS, tmp_path, REQS)
    b = next(x for x in col["bundles"] if x["test_code"] == "TC-001")
    assert b["criterion_statement"] == "The system generates a BRD document"
    assert b["observed"]["ui_output"] == "Your BRD is ready\nDocuments generated successfully"
    assert "Documents generated successfully" in b["observed"]["messages"]
    doc = b["observed"]["documents"][0]
    assert doc["readable"] and doc["format"] == "docx" and "Problem statement" in doc["text_excerpt"]
    assert (tmp_path / doc["text_file"]).read_text(encoding="utf-8").startswith("Business Requirements Document")
    assert [i["synthetic"] for i in b["inputs"]] == [False, True]
    types = {a["type"] for a in b["artifacts"]}
    assert {"screenshot", "page_text", "page_text_before", "download", "document_text"} <= types
    shot = next(a for a in b["artifacts"] if a["type"] == "screenshot")
    assert shot["sha256"] == hashlib.sha256(b"\x89PNG fake").hexdigest()
    assert b["completeness"] == {"output_captured": True, "expected_output": "document", "missing": []}
    assert b["observations"]["api_calls"] == 2 and b["observations"]["api_failures"][0]["status"] == 500
    assert b["observations"]["slowest_call"]["path"] == "/api/generate"
    manifest = json.loads((tmp_path / "evidence_manifest.json").read_text(encoding="utf-8"))
    assert manifest["bundles"][0]["digest"] == b["digest"]


def test_missing_document_is_stated_and_unmappable_test_recorded(tmp_path):
    col = collect_evidence({"results": [_result(tmp_path, with_download=False)]}, PLAN, TESTS, tmp_path, REQS)
    b1 = next(x for x in col["bundles"] if x["test_code"] == "TC-001")
    assert b1["completeness"]["output_captured"] is False
    assert "expected a downloaded document but none was produced" in b1["completeness"]["missing"]
    b2 = next(x for x in col["bundles"] if x["test_code"] == "TC-002")
    assert b2["execution_status"] == "not_executed" and "latency claim" in b2["error"]
    assert col["counts"]["not_executed"] == 1 and col["by_criterion"] == {"AC-01.1": ["EV-TC-001"], "AC-02.1": ["EV-TC-002"]}


def test_blocked_plan_yields_not_executed_bundle_per_test(tmp_path):
    col = collect_evidence(None, {"status": "blocked", "reason": "behind a login", "plans": []}, TESTS, tmp_path)
    assert [b["execution_status"] for b in col["bundles"]] == ["not_executed", "not_executed"]
    assert all("behind a login" in b["error"] for b in col["bundles"])


def test_reordered_page_counts_as_output():
    from engines.runtime.evidence import _appeared

    text, kind = _appeared("Products\nBackpack $29.99\nOnesie $7.99\nJacket $49.99",
                           "Products\nOnesie $7.99\nBackpack $29.99\nJacket $49.99")
    assert kind == "reordered" and "Onesie $7.99" in text and "Products" not in text
    assert _appeared("a\nb", "a\nb") == ("", "none")
    text, kind = _appeared("Form\nGenerate", "Form\nGenerate\nYour BRD is ready")
    assert (text, kind) == ("Your BRD is ready", "added")
